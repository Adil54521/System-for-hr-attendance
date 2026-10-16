from django.contrib import admin
from django.contrib.admin.options import csrf_protect_m

from .models import Teacher, VisitLog, MessageHistory
import openpyxl
from django.http import HttpResponse
from datetime import datetime, timedelta, date, time
from .models import WebAppAccessLog
from .models import SuspiciousAccessLog
from .models import AbiturientsAccessCount
from rangefilter.filters import DateRangeFilterBuilder
from myapp.tabel import ExportToTabel
from myapp.export_tabel import ExportToTabelByDate 
from django.shortcuts import render
#import google.generativeai as genai
from django.conf import settings
from django.db.models import Count, Q
from openpyxl.utils import get_column_letter
from django.db.models import OuterRef, Subquery
import json
from django.contrib.contenttypes.models import ContentType
class LocationListFilter(admin.SimpleListFilter):
    title = "Локация"
    parameter_name = "location"

    def lookups(self, request, model_admin):
        # Получаем уникальные локации из VisitLog
        locations = VisitLog.objects.values_list('location', flat=True).distinct()
        return [(loc, loc) for loc in locations if loc]

    def queryset(self, request, queryset):
        if self.value():
            return queryset.filter(location=self.value())
        return queryset

class VisitDateRangeFilter(admin.SimpleListFilter):
    title = "Диапазон дат визитов"
    parameter_name = "visit_date"
    template = "admin/date_range_filter.html"  # кастомный шаблон с двумя input type="date"

    def lookups(self, request, model_admin):
        return ()  # пусто, значения вводятся вручную

    def queryset(self, request, queryset):
        start = request.GET.get("start_date")
        end = request.GET.get("end_date")
        if start and end:
            # фильтруем Teacher через связь с VisitLog
            return queryset.filter(
                visitlog__date__range=[parse_date(start), parse_date(end)]
            ).distinct()
        return queryset

class MultiFioListFilter(admin.SimpleListFilter):
    title = "ФИО"
    parameter_name = "fio"
    template = "admin/multi_filter.html"

    def lookups(self, request, model_admin):
        return [(t.fio, t.fio) for t in Teacher.objects.all().order_by("fio")]

    def queryset(self, request, queryset):
        values = request.GET.getlist(self.parameter_name)
        if values:
            return queryset.filter(fio__in=values)
        return queryset

    def value(self):
        """Возвращаем список выбранных значений, а не одно"""
        return self.used_parameters.get(self.parameter_name, [])

@admin.action(description="📥 Скачать выбранные записи в Excel")
def export_teachers_to_excel(modeladmin, request, queryset):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сотрудники"

    # Заголовки
    headers = ["ФИО", "Должность", "Подразделение", "Телефон"]
    ws.append(headers)

    # Данные
    for teacher in queryset:
        ws.append([
            teacher.fio,
            teacher.job_title,
            teacher.department,
            teacher.phone
        ])

    # Ответ
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename=teachers.xlsx'
    wb.save(response)
    return response

# Регистрация модели Teacher с actions
@admin.register(Teacher)
class TeacherAdmin(admin.ModelAdmin):
    list_display = ("fio", "job_title", "department", "phone", "telegram_id")
    search_fields = ("fio", "job_title", "department", "phone", "phone_clean")
    list_filter = (MultiFioListFilter, VisitDateRangeFilter)
    actions = ["export_teach_to_excel", "export_to_excel", "some_func"]
    # change_list_template = "admin/send_message.html"

    def export_teach_to_excel(modeladmin, request, queryset):
        # выбираем все VisitLog для выбранных учителей
        visitlogs = VisitLog.objects.filter(teacher__in=queryset)
        
        # инициализация ExportToTabelByDate с request и queryset
        exporter = ExportToTabelByDate(request, queryset=visitlogs)
        
        exporter.teacher = queryset
    
        exporter.excelOutPut()
        exporter.placeDate()
        exporter.placeNameJob(request)
        
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"Табель_{datetime.today().strftime('%Y-%m-%d')}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
    
        exporter.workbook.save(response)
        return response

    export_teach_to_excel.short_description = "⬇️ Сформировать табель в Excel"

    def export_to_excel(modeladmin, request, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Данные"

        # Заголовки
        headers = [field.name for field in modeladmin.model._meta.fields]
        ws.append(headers)

        # Данные
        for obj in queryset:
            row = [getattr(obj, field) for field in headers]
            ws.append(row)

        for i, column in enumerate(ws.columns, start=1):
            max_length = max(len(str(cell.value)) if cell.value else 0 for cell in column)
            ws.column_dimensions[get_column_letter(i)].width = max_length + 2

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response['Content-Disposition'] = 'attachment; filename="export.xlsx"'
        wb.save(response)
        return response

    export_to_excel.short_description = "⬇️ Сформировать список сотрудников в Excel"

    def some_func(modeladmin, request, queryset):
        selected_teachers = list(queryset)

        selected_phones = [
            teacher.phone
            for teacher in selected_teachers
            if teacher.phone
        ]

        # Подзапрос на последнее сообщение сотрудника.
        latest_history_item = (
            MessageHistory.objects
            .filter(
                recipient_id=OuterRef("recipient_id")
            )
            .order_by(
                "-created_at",
                "-pk"
            )
            .values("pk")[:1]
        )

        # Получаем только одно последнее сообщение
        # для каждого сотрудника.
        history = (
            MessageHistory.objects
            .select_related("recipient")
            .filter(
                pk=Subquery(latest_history_item)
            )
            .order_by(
                "-created_at",
                "-pk"
            )
        )

        conversations = [
            {
                "history_id": history_item.pk,
                "recipient": history_item.recipient,
                "last_subject": history_item.subject,
                "last_message": history_item.message,
                "updated_at": history_item.created_at,
                "read_at": history_item.read_at,
                "is_read": (
                        history_item.read_at is not None
                ),
            }
            for history_item in history
        ]

        return render(
            request,
            "admin/send_message.html",
            {
                # Справа отображаются только сотрудники,
                # выбранные через Django Admin.
                "q_obj": selected_teachers,
                "selected_phones": selected_phones,
                "selected_count": len(
                    selected_teachers
                ),

                # Слева отображаются диалоги всех сотрудников,
                # которым ранее отправлялись сообщения.
                "conversations": conversations,
            }
        )

    some_func.short_description = ("Сообщение для сотрудников")

@admin.register(VisitLog)
class VisitLogAdmin(admin.ModelAdmin):
    list_display = ('teacher', 'date', 'location', 'time_in', 'time_out', 'attendance', 'check_gal')
    list_filter = (('date', DateRangeFilterBuilder()), LocationListFilter)
    search_fields = ['teacher__fio', 'teacher__job_title', 'teacher__department']
    actions = ["delete_selected","show_absent", "export_visitlog_to_excel"]

    change_list_template = "admin/visitlog_changelist.html"

    def show_absent(self, request, queryset):
        today = date.today()
        all_teachers = set(Teacher.objects.values_list("id", flat=True))
        present_today = set(VisitLog.objects.filter(date=today).values_list("teacher_id", flat=True))
        absent_ids = all_teachers - present_today

        # Получим список отсутствующих
        absent_teachers = Teacher.objects.filter(id__in=absent_ids)

        # Сообщим пользователю
        self.message_user(request, f"📋 Отсутствующие сегодня ({today}):")
        for teacher in absent_teachers:
            self.message_user(request, f"— {teacher.fio} | {teacher.job_title}")

    show_absent.short_description = "🚫 Показать отсутствующих сегодня"
    
    def export_visitlog_to_excel(modeladmin, request, queryset):
        exporter = ExportToTabel(request)
        exporter.excelOutPut()
        exporter.placeDate()
        exporter.placeNameJob(request)
        
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"Табель_{datetime.today().strftime('%Y-%m-%d')}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        
        exporter.workbook.save(response)
        return response

    export_visitlog_to_excel.short_description = "⬇️ Сформировать в табель в формате Excel"

@admin.register(WebAppAccessLog)
class WebAppAccessLogAdmin(admin.ModelAdmin):
    list_display = ("timestamp", "ip", "success", "campus", "reason", "accuracy", "phone", "phone_matched")
    list_filter = ("success", "campus")
    search_fields = ("ip", "user_agent", "reason", "phone", "phone_matched")

@admin.register(SuspiciousAccessLog)
class SuspiciousAccessLogAdmin(admin.ModelAdmin):
    list_display = ("ip_address", "timestamp", "reason", "blocked")
    list_filter = ("blocked", "reason")
    search_fields = ("ip_address", "user_agent")

@admin.register(AbiturientsAccessCount)
class AbiturientsAccessCount(admin.ModelAdmin):
    list_display = ("ip", "user_agent")
    list_filter = (('date', DateRangeFilterBuilder()), "ip", "user_agent")
    search_fields = ("ip", "user_agent")
